"""Local document ingestion and immutable, fingerprinted search generations."""

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from collections import Counter
from pathlib import Path

import faiss
import numpy as np
from filelock import FileLock, Timeout

from eve_ah.config import Settings
from eve_ah.models import EveError, Source

FORMAT_VERSION = 4
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}
LEXICAL_STOPWORDS = {
    "a",
    "an",
    "and",
    "are",
    "can",
    "do",
    "for",
    "how",
    "i",
    "in",
    "is",
    "it",
    "me",
    "of",
    "on",
    "the",
    "to",
    "what",
    "when",
    "with",
    "you",
}


def _hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


class LocalEmbedder:
    """Lazy model loading keeps CLI help and process health lightweight."""

    def __init__(self, settings: Settings):
        from sentence_transformers import SentenceTransformer

        self.model = SentenceTransformer(settings.embed_model)
        self.tokenizer = self.model.tokenizer
        self.limit = int(self.model.max_seq_length) - self.tokenizer.num_special_tokens_to_add()
        revision = getattr(self.model[0].auto_model.config, "_commit_hash", None)
        self.fingerprint = {
            "model": settings.embed_model,
            "revision": revision,
            "dimension": self.model.get_embedding_dimension(),
            "tokenizer": _hash(self.tokenizer.backend_tokenizer.to_str().encode()),
            "max_tokens": self.limit,
            "normalized": True,
        }
        if settings.bot.chunk_tokens > self.limit:
            raise EveError(
                f"chunk_tokens exceeds this embedding model's limit ({self.limit}); update bot.yaml.",
                code="invalid_chunking",
            )

    def chunk(self, text: str, tokens: int, overlap: int) -> list[str]:
        if tokens > self.limit or overlap >= tokens:
            raise EveError("Invalid tokenizer chunk settings.", code="invalid_chunking")
        encoded = self.tokenizer(text, add_special_tokens=False, return_offsets_mapping=True)
        offsets = encoded["offset_mapping"]
        pieces = []
        start = 0
        while start < len(offsets):
            end = min(start + tokens, len(offsets))
            # Prefer a paragraph boundary in the latter half without exceeding token budget.
            if end < len(offsets):
                for boundary in range(end - 1, start + tokens // 2, -1):
                    gap = text[offsets[boundary - 1][1] : offsets[boundary][0]]
                    if "\n" in gap:
                        end = boundary
                        break
            piece = text[offsets[start][0] : offsets[end - 1][1]].strip()
            if piece:
                pieces.append(piece)
            if end == len(offsets):
                break
            start = max(start + 1, end - overlap)
        return pieces

    def embed(self, texts: list[str]) -> np.ndarray:
        return np.asarray(
            self.model.encode(texts, normalize_embeddings=True, show_progress_bar=False),
            dtype=np.float32,
        )


def _extract(path: Path, max_pages: int, max_bytes: int) -> list[dict]:
    """Runs in a child process; the parent applies a hard extraction timeout."""
    suffix = path.suffix.lower()
    pages = []
    if suffix in {".txt", ".md"}:
        pages = [{"page": None, "text": path.read_text(encoding="utf-8")}]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF is unsupported")
        if len(reader.pages) > max_pages:
            raise ValueError("PDF exceeds page limit")
        for i, page in enumerate(reader.pages):
            pages.append({"page": i + 1, "text": page.extract_text() or ""})
            if sum(len(p["text"].encode()) for p in pages) > max_bytes:
                raise ValueError("Extracted text exceeds size limit")
    elif suffix == ".docx":
        from docx import Document

        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > max_bytes * 5:
                raise ValueError("Expanded document exceeds size limit")
        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        pages = [{"page": None, "text": text}]
    if sum(len(p["text"].encode()) for p in pages) > max_bytes:
        raise ValueError("Extracted text exceeds size limit")
    pages = [p for p in pages if p["text"].strip()]
    if not pages:
        raise ValueError("No extractable text; image-only PDFs require OCR")
    return pages


class Knowledge:
    def __init__(self, settings: Settings, embedder=None):
        self.settings = settings
        self._embedder = embedder
        self.index = None
        self.records = []
        self.catalog = {"topics": [], "suggested_questions": []}
        self.lexical = {"documents": [], "doc_freq": {}, "postings": {}, "avg_length": 0}
        self.generation = None

    @property
    def embedder(self):
        if self._embedder is None:
            self._embedder = LocalEmbedder(self.settings)
        return self._embedder

    def _fingerprint(self):
        return {
            "format": FORMAT_VERSION,
            "embedding": self.embedder.fingerprint,
            "chunk_tokens": self.settings.bot.chunk_tokens,
            "chunk_overlap": self.settings.bot.chunk_overlap,
            "catalog": {
                "topics": [topic.model_dump() for topic in self.settings.bot.topics],
                "suggested_questions": self.settings.bot.suggested_questions,
            },
        }

    def _read_generation(self):
        try:
            pointer = json.loads((self.settings.storage_dir / "active.json").read_text())
            generation = pointer["generation"]
            if not isinstance(generation, str) or len(generation) != 32:
                raise ValueError("Invalid generation")
            int(generation, 16)
            folder = self.settings.storage_dir / "generations" / generation
            manifest_bytes = (folder / "manifest.json").read_bytes()
            if _hash(manifest_bytes) != pointer["manifest_sha256"]:
                raise ValueError("Manifest checksum mismatch")
            manifest = json.loads(manifest_bytes)
            if manifest["fingerprint"] != self._fingerprint():
                raise EveError(
                    "Index is incompatible with configuration; run eve-ah ingest.",
                    code="index_incompatible",
                )
            record_bytes = (folder / "chunks.json").read_bytes()
            if _hash(record_bytes) != manifest["chunks_sha256"]:
                raise ValueError("Chunk checksum mismatch")
            index_bytes = (folder / "index.faiss").read_bytes()
            if _hash(index_bytes) != manifest["index_sha256"]:
                raise ValueError("Index checksum mismatch")
            records = json.loads(record_bytes)
            catalog_bytes = (folder / "catalog.json").read_bytes()
            if _hash(catalog_bytes) != manifest["catalog_sha256"]:
                raise ValueError("Catalog checksum mismatch")
            catalog = json.loads(catalog_bytes)
            lexical_bytes = (folder / "lexical.json").read_bytes()
            if _hash(lexical_bytes) != manifest["lexical_sha256"]:
                raise ValueError("Lexical checksum mismatch")
            lexical = json.loads(lexical_bytes)
            index = faiss.deserialize_index(np.frombuffer(index_bytes, dtype=np.uint8).copy())
            if index.ntotal != len(records) or index.d != self.embedder.fingerprint["dimension"]:
                raise ValueError("Index dimensions mismatch")
            if not isinstance(catalog.get("topics"), list) or not isinstance(
                catalog.get("suggested_questions"), list
            ):
                raise ValueError("Invalid catalog")
            if len(lexical.get("documents", [])) != len(records):
                raise ValueError("Lexical index size mismatch")
            if not isinstance(lexical.get("postings"), dict):
                raise ValueError("Invalid lexical postings")
            return generation, index, records, catalog, lexical
        except EveError:
            raise
        except (OSError, ValueError, KeyError, TypeError, RuntimeError) as exc:
            raise EveError(
                "Index unavailable or damaged; run eve-ah ingest.", code="index_unavailable"
            ) from exc

    def load(self):
        # Clear stale in-process data even when reload fails.
        self.index, self.records, self.generation = None, [], None
        self.catalog = {"topics": [], "suggested_questions": []}
        self.lexical = {"documents": [], "doc_freq": {}, "postings": {}, "avg_length": 0}
        generation, index, records, catalog, lexical = self._read_generation()
        if not records:
            raise EveError(
                "No documents indexed; add documents and run eve-ah ingest.", code="empty_index"
            )
        self.generation, self.index, self.records = generation, index, records
        self.catalog, self.lexical = catalog, lexical

    def overview(self):
        return self.catalog

    def _files(self):
        root = self.settings.data_dir
        if root.is_symlink() or not root.is_dir():
            raise EveError(
                "The data directory must exist and must not be a symlink.", code="invalid_data"
            )
        files, skipped, total = [], [], 0

        def traversal_error(error):
            raise EveError(
                "Cannot read the complete data directory; check permissions.", code="invalid_data"
            ) from error

        for base, directories, names in os.walk(root, followlinks=False, onerror=traversal_error):
            for name in directories + names:
                if (Path(base) / name).is_symlink():
                    raise EveError(
                        "Symlinks are not allowed in the data directory.", code="invalid_data"
                    )
            for name in sorted(names):
                path = Path(base) / name
                relative = path.relative_to(root).as_posix()
                if path.suffix.lower() not in SUPPORTED:
                    skipped.append(relative)
                    continue
                size = path.stat().st_size
                total += size
                if size > self.settings.max_file_bytes or total > self.settings.max_corpus_bytes:
                    raise EveError(
                        "Document or corpus exceeds configured size limit.", code="size_limit"
                    )
                files.append((relative, path))
                if len(files) > 10000:
                    raise EveError("Corpus exceeds 10000 documents.", code="size_limit")
        return sorted(files), skipped

    def _parse(self, path):
        # Child process isolates slow parsers and can be killed on timeout.
        try:
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "eve_ah.knowledge",
                    "--extract",
                    str(path),
                    str(self.settings.max_pages),
                    str(self.settings.max_file_bytes),
                ],
                capture_output=True,
                text=True,
                timeout=30,
                check=True,
                env={**os.environ, "PYTHONPATH": os.pathsep.join(sys.path)},
            )
            return json.loads(result.stdout)
        except (subprocess.SubprocessError, ValueError) as exc:
            raise EveError(
                f"Could not extract {path.name}; check format, size, page limits, or OCR requirements.",
                code="extraction_failed",
            ) from exc

    def ingest(self):
        self.settings.storage_dir.mkdir(parents=True, exist_ok=True)
        try:
            with FileLock(str(self.settings.storage_dir / "writer.lock"), timeout=0):
                return self._ingest_locked()
        except Timeout as exc:
            raise EveError("An ingestion is already running.", code="ingestion_busy") from exc

    def _ingest_locked(self):
        files, skipped = self._files()
        fingerprint = self._fingerprint()
        previous = {}
        try:
            _, old_index, old_records, _, _ = self._read_generation()
            vectors = old_index.reconstruct_n(0, old_index.ntotal)
            previous = {r["text_hash"]: vectors[i] for i, r in enumerate(old_records)}
        except EveError:
            pass  # A complete rebuild repairs an absent, incompatible, or damaged index.
        records, hashes, extracted_bytes = [], {}, 0
        for name, path in files:
            hashes[name] = _hash(path.read_bytes())
            for page in self._parse(path):
                extracted_bytes += len(page["text"].encode())
                if extracted_bytes > self.settings.max_corpus_bytes:
                    raise EveError(
                        "Extracted corpus exceeds configured size limit.", code="size_limit"
                    )
                chunks = self.embedder.chunk(
                    page["text"],
                    self.settings.bot.chunk_tokens,
                    self.settings.bot.chunk_overlap,
                )
                page_heading = next(
                    (
                        match.group(1).strip()
                        for match in re.finditer(r"(?m)^#{1,3}\s+(.+)$", page["text"])
                    ),
                    None,
                )
                page_description = next(
                    (
                        re.sub(r"\s+", " ", paragraph).strip()[:300]
                        for paragraph in re.split(r"\n\s*\n", page["text"])
                        if paragraph.strip() and not paragraph.lstrip().startswith("#")
                    ),
                    "",
                )
                for i, text in enumerate(chunks):
                    text_hash = _hash(text.encode())
                    source_id = "c-" + _hash(f"{name}:{page['page']}:{i}:{text_hash}".encode())[:16]
                    records.append(
                        {
                            "id": source_id,
                            "name": name,
                            "page": page["page"],
                            "text": text,
                            "text_hash": text_hash,
                            "heading": page_heading,
                            "description": page_description,
                        }
                    )
        dimension = self.embedder.fingerprint["dimension"]
        missing = {r["text_hash"]: r["text"] for r in records if r["text_hash"] not in previous}
        if missing:
            keys = list(missing)
            for start in range(0, len(keys), 64):
                batch = keys[start : start + 64]
                encoded = self.embedder.embed([missing[k] for k in batch])
                if encoded.shape != (len(batch), dimension) or not np.isfinite(encoded).all():
                    raise EveError(
                        "Embedding model returned invalid vectors.", code="embedding_failed"
                    )
                previous.update(zip(batch, encoded))
        index = faiss.IndexFlatIP(dimension)
        if records:
            vectors = np.asarray([previous[r["text_hash"]] for r in records], dtype=np.float32)
            faiss.normalize_L2(vectors)
            index.add(vectors)
        catalog = self._build_catalog(records)
        lexical = self._build_lexical(records)
        generation = uuid.uuid4().hex
        generations = self.settings.storage_dir / "generations"
        generations.mkdir(exist_ok=True)
        draft = Path(tempfile.mkdtemp(prefix=".build-", dir=generations))
        try:
            chunks_bytes = json.dumps(records, ensure_ascii=False).encode()
            catalog_bytes = json.dumps(catalog, ensure_ascii=False, sort_keys=True).encode()
            lexical_bytes = json.dumps(lexical, ensure_ascii=False, sort_keys=True).encode()
            index_bytes = faiss.serialize_index(index).tobytes()
            (draft / "chunks.json").write_bytes(chunks_bytes)
            (draft / "catalog.json").write_bytes(catalog_bytes)
            (draft / "lexical.json").write_bytes(lexical_bytes)
            (draft / "index.faiss").write_bytes(index_bytes)
            manifest = {
                "fingerprint": fingerprint,
                "files": hashes,
                "count": len(records),
                "chunks_sha256": _hash(chunks_bytes),
                "catalog_sha256": _hash(catalog_bytes),
                "lexical_sha256": _hash(lexical_bytes),
                "index_sha256": _hash(index_bytes),
            }
            manifest_bytes = json.dumps(manifest, sort_keys=True).encode()
            (draft / "manifest.json").write_bytes(manifest_bytes)
            draft.rename(generations / generation)
            pointer = {"generation": generation, "manifest_sha256": _hash(manifest_bytes)}
            pending = self.settings.storage_dir / f".active-{generation}.json"
            pending.write_text(json.dumps(pointer))
            os.replace(pending, self.settings.storage_dir / "active.json")
        finally:
            if draft.exists():
                shutil.rmtree(draft)
        return {
            "accepted": [name for name, _ in files],
            "skipped": skipped,
            "chunks": len(records),
            "generation": generation,
            "embedded": len(missing),
            "topics": len(catalog["topics"]),
        }

    def _build_catalog(self, records: list[dict]) -> dict:
        by_source: dict[str, list[dict]] = {}
        for record in records:
            by_source.setdefault(record["name"], []).append(record)

        if self.settings.bot.topics:
            topics = []
            for configured in self.settings.bot.topics:
                unknown = sorted(set(configured.sources).difference(by_source))
                if unknown:
                    raise EveError(
                        f"Topic {configured.name!r} references unknown source: {unknown[0]}",
                        code="invalid_topics",
                    )
                source_ids = [
                    record["id"] for source in configured.sources for record in by_source[source]
                ]
                topics.append(
                    {
                        "name": configured.name,
                        "description": configured.description,
                        "sources": configured.sources,
                        "source_ids": source_ids,
                    }
                )
        else:
            topics = []
            for source, source_records in sorted(by_source.items()):
                heading = next(
                    (record.get("heading") for record in source_records if record.get("heading")),
                    None,
                )
                name = heading or Path(source).stem.replace("-", " ").replace("_", " ").title()
                description = next(
                    (
                        record.get("description", "")
                        for record in source_records
                        if record.get("description")
                    ),
                    "",
                )
                topics.append(
                    {
                        "name": name,
                        "description": description,
                        "sources": [source],
                        "source_ids": [record["id"] for record in source_records],
                    }
                )
        return {
            "topics": topics,
            "suggested_questions": list(self.settings.bot.suggested_questions),
        }

    @staticmethod
    def _lexical_tokens(text: str) -> list[str]:
        return [
            token
            for token in re.findall(r"(?u)\b[\w]+(?:[-.][\w]+)*\b", text.lower())
            if len(token) > 1 and token not in LEXICAL_STOPWORDS
        ]

    def _build_lexical(self, records: list[dict]) -> dict:
        documents = []
        document_frequency: Counter[str] = Counter()
        postings: dict[str, list[list[int]]] = {}
        for index, record in enumerate(records):
            counts = Counter(self._lexical_tokens(record["text"]))
            documents.append({"length": sum(counts.values())})
            document_frequency.update(counts.keys())
            for term, frequency in counts.items():
                postings.setdefault(term, []).append([index, frequency])
        average = (
            sum(document["length"] for document in documents) / len(documents) if documents else 0
        )
        return {
            "documents": documents,
            "doc_freq": dict(document_frequency),
            "postings": postings,
            "avg_length": average,
        }

    def _lexical_scores(self, query: str) -> dict[int, float]:
        tokens = list(dict.fromkeys(self._lexical_tokens(query)))
        if not tokens or not self.records:
            return {}
        total = len(self.records)
        average = self.lexical["avg_length"] or 1
        term_postings = {token: dict(self.lexical["postings"].get(token, [])) for token in tokens}
        candidates = {index for postings in term_postings.values() for index in postings}
        scores: dict[int, float] = {}
        for index in candidates:
            document = self.lexical["documents"][index]
            score = 0.0
            matched = []
            for token in tokens:
                frequency = term_postings[token].get(index, 0)
                if not frequency:
                    continue
                matched.append(token)
                document_frequency = self.lexical["doc_freq"].get(token, 0)
                inverse = np.log(
                    1 + (total - document_frequency + 0.5) / (document_frequency + 0.5)
                )
                denominator = frequency + 1.5 * (1 - 0.75 + 0.75 * document["length"] / average)
                score += inverse * frequency * 2.5 / denominator
            has_unique_identifier_match = any(
                ("-" in token or any(char.isdigit() for char in token))
                and self.lexical["doc_freq"].get(token, total) == 1
                for token in matched
            )
            if score > 0 and (len(matched) >= 2 or has_unique_identifier_match):
                scores[index] = float(score)
        return scores

    def search(self, query: str, mode: str | None = None) -> list[Source]:
        if self.index is None:
            raise EveError("Index not loaded; restart after ingestion.", code="index_unavailable")
        mode = mode or self.settings.bot.retrieval_mode
        if mode not in {"dense", "hybrid"}:
            raise ValueError("Unknown retrieval mode")
        vector = np.asarray(self.embedder.embed([query]), dtype=np.float32)
        faiss.normalize_L2(vector)
        scores, indices = self.index.search(
            vector, min(self.index.ntotal, self.settings.bot.top_k * 6)
        )
        dense_scores = {
            int(index): float(score) for score, index in zip(scores[0], indices[0]) if index >= 0
        }
        if mode == "dense":
            ranked = [
                (int(index), float(score))
                for score, index in zip(scores[0], indices[0])
                if index >= 0 and score >= self.settings.bot.min_score
            ]
        else:
            lexical_scores = self._lexical_scores(query)
            dense_ranks = {
                int(index): rank for rank, index in enumerate(indices[0], 1) if index >= 0
            }
            lexical_ranks = {
                index: rank
                for rank, (index, _) in enumerate(
                    sorted(lexical_scores.items(), key=lambda item: (-item[1], item[0])), 1
                )
            }
            candidates = set(dense_ranks).union(lexical_ranks)
            fused = {}
            for index in candidates:
                if (
                    dense_scores.get(index, -1) < self.settings.bot.min_score
                    and index not in lexical_ranks
                ):
                    continue
                score = 0.0
                if index in dense_ranks:
                    score += 1 / (self.settings.bot.rrf_k + dense_ranks[index])
                if index in lexical_ranks:
                    score += 1 / (self.settings.bot.rrf_k + lexical_ranks[index])
                fused[index] = score
            ranked = sorted(fused.items(), key=lambda item: (-item[1], item[0]))
        results, seen, remaining = [], set(), self.settings.bot.context_chars
        for index, score in ranked:
            record = self.records[index]
            if record["text_hash"] in seen:
                continue
            text = record["text"]
            if len(text) > remaining:
                continue
            seen.add(record["text_hash"])
            remaining -= len(text)
            results.append(
                Source(
                    id=record["id"],
                    name=record["name"],
                    page=record["page"],
                    text=text,
                    score=float(score),
                )
            )
            if len(results) >= self.settings.bot.top_k:
                break
        return results


if __name__ == "__main__":
    if len(sys.argv) == 5 and sys.argv[1] == "--extract":
        try:
            print(json.dumps(_extract(Path(sys.argv[2]), int(sys.argv[3]), int(sys.argv[4]))))
        except Exception:
            sys.exit(1)
    else:
        sys.exit(2)
